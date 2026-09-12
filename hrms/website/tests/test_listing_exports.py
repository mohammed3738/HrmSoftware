"""
Regression tests for the Branches, Companies, and Employees Offboarding
list pages' Export buttons -- Branches and Companies both had "Export as
Excel" dropdown links wired to href="javascript:void(0);" (dead placeholders
copied from the page template), and Offboarding's Export button called
window.print() instead of producing an actual spreadsheet. All three now
have a dedicated export endpoint and a real link/onclick pointing at it.

Run with: python manage.py test website.tests.test_listing_exports
"""
import io
from datetime import date

import openpyxl
from django.contrib.auth.models import User
from django.test import TestCase, Client
from django.urls import reverse

from website.models import Branch, Company, Employee, Offboarding


class BranchExportTest(TestCase):
    def setUp(self):
        Branch.objects.create(branch_name="Head Office", branch_address="123 Main St", is_active=True)
        Branch.objects.create(branch_name="Old Branch", branch_address="456 Side St", is_active=False)
        self.user = User.objects.create_superuser("branch_admin", "branch@test.com", "pass12345")
        self.client = Client()
        self.client.force_login(self.user, backend="django.contrib.auth.backends.ModelBackend")

    def test_export_url_is_reachable_and_returns_an_excel_file(self):
        resp = self.client.get(reverse("branch-export-excel"))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(
            resp["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def test_default_export_only_includes_active_branches(self):
        resp = self.client.get(reverse("branch-export-excel"))
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(rows), 1)
        header = [c.value for c in ws[1]]
        row = dict(zip(header, rows[0]))
        self.assertEqual(row["Branch Name"], "Head Office")
        self.assertEqual(row["Status"], "Active")

    def test_show_all_includes_the_archived_branch_too(self):
        resp = self.client.get(reverse("branch-export-excel"), {"show": "all"})
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(rows), 2)

    def test_the_page_links_to_the_dedicated_export_endpoint_not_a_dead_link(self):
        resp = self.client.get(reverse("create-branch"))
        self.assertContains(resp, "downloadBranches")
        self.assertContains(resp, reverse("branch-export-excel"))


class CompanyExportTest(TestCase):
    def setUp(self):
        Company.objects.create(
            short_name="ZC", name="Zeta Corp", phone="1234567890", email="zc@test.com",
            address="1 Zeta Rd", status="active",
        )
        Company.objects.create(
            short_name="OC", name="Old Corp", phone="0987654321", email="oc@test.com",
            address="1 Old Rd", status="inactive",
        )
        self.user = User.objects.create_superuser("company_admin", "company@test.com", "pass12345")
        self.client = Client()
        self.client.force_login(self.user, backend="django.contrib.auth.backends.ModelBackend")

    def test_export_url_is_reachable_and_returns_an_excel_file(self):
        resp = self.client.get(reverse("company-export-excel"))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(
            resp["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def test_default_export_only_includes_active_companies(self):
        resp = self.client.get(reverse("company-export-excel"))
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(rows), 1)
        header = [c.value for c in ws[1]]
        row = dict(zip(header, rows[0]))
        self.assertEqual(row["Company Name"], "Zeta Corp")
        self.assertEqual(row["Status"], "active")

    def test_show_all_includes_the_inactive_company_too(self):
        resp = self.client.get(reverse("company-export-excel"), {"show": "all"})
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(rows), 2)

    def test_the_page_links_to_the_dedicated_export_endpoint_not_a_dead_link(self):
        resp = self.client.get(reverse("create-company"))
        self.assertContains(resp, "downloadCompanies")
        self.assertContains(resp, reverse("company-export-excel"))


class OffboardingExportTest(TestCase):
    def setUp(self):
        self.employee = Employee.objects.create(
            salutation="Mr", first_name="Off", last_name="Board",
            father_name="Robert Doe", gender="Male", blood_group="O+",
            date_of_birth=date(1990, 1, 1), place_of_birth="Test City",
            personal_email="offboard@test.com", present_address="123 Test St",
            permanent_address="123 Test St", personal_mobile="1234567890",
            employee_code="OFFEXP1", designation="Developer", department="IT",
            date_of_joining=date(2020, 1, 1), location="Test Location",
            pan_no="ABCDE1234F", aadhar_no="123456789012",
            name_as_per_bank="Off Board", salary_account_number="1234567890",
            ifsc_code="TEST0001234", emergency_contact_name1="Jane Doe",
            emergency_contact_relation1="Spouse", emergency_contact_mobile1="0987654321",
            status="Left",
        )
        Offboarding.objects.create(
            employee=self.employee,
            date_of_resignation=date(2025, 1, 1),
            date_of_relieving=date(2025, 2, 1),
            is_active=True,
        )
        self.user = User.objects.create_superuser("offboard_admin", "offboard@test.com", "pass12345")
        self.client = Client()
        self.client.force_login(self.user, backend="django.contrib.auth.backends.ModelBackend")

    def test_export_url_is_reachable_and_returns_an_excel_file(self):
        resp = self.client.get(reverse("offboarding-export-excel"))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(
            resp["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def test_export_contains_the_offboarding_row(self):
        resp = self.client.get(reverse("offboarding-export-excel"))
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        header = [c.value for c in ws[1]]
        row = dict(zip(header, [c.value for c in ws[2]]))
        self.assertEqual(row["Employee Code"], "OFFEXP1")
        self.assertEqual(row["Status"], "Active")

    def test_the_page_links_to_the_dedicated_export_endpoint_not_window_print(self):
        resp = self.client.get(reverse("offboarding-list"))
        self.assertContains(resp, "downloadOffboarding")
        self.assertContains(resp, reverse("offboarding-export-excel"))
        self.assertNotContains(resp, "onclick=\"window.print()\"")
