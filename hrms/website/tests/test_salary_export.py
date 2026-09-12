"""
Regression test: the Salary Master and Salary Increments pages' Export
buttons were wired to the wrong thing -- Salary Master's button called the
shared downloadEmployees() helper (defined once in base2.html for the
Employee list page, with no override of its own), so it silently downloaded
the employee list instead of salary data. Salary Increments' button called
exportIncrements(), which was never defined anywhere, so it did nothing.

Both pages now have their own dedicated export endpoint and a local JS
override that points at it, matching the working pattern already used by
Salary History's export.

Run with: python manage.py test website.tests.test_salary_export
"""
import io
from datetime import date

import openpyxl
from django.contrib.auth.models import User
from django.test import TestCase, Client
from django.urls import reverse

from website.models import Company, Employee, SalaryMaster, SalaryIncrement


def make_employee(code, first="Sal", last="Ary"):
    return Employee.objects.create(
        salutation="Mr", first_name=first, last_name=last,
        father_name="Robert Doe", gender="Male", blood_group="O+",
        date_of_birth=date(1990, 1, 1), place_of_birth="Test City",
        personal_email=f"{code.lower()}@test.com", present_address="123 Test St",
        permanent_address="123 Test St", personal_mobile="1234567890",
        employee_code=code, designation="Developer", department="IT",
        date_of_joining=date(2020, 1, 1), location="Test Location",
        pan_no="ABCDE1234F", aadhar_no="123456789012",
        name_as_per_bank="Sal Ary", salary_account_number="1234567890",
        ifsc_code="TEST0001234", emergency_contact_name1="Jane Doe",
        emergency_contact_relation1="Spouse", emergency_contact_mobile1="0987654321",
        status="Active",
    )


class SalaryMasterExportTest(TestCase):
    def setUp(self):
        self.employee = make_employee("SALEXP1")
        SalaryMaster.objects.create(
            employee=self.employee,
            gross_ctc_pm=50000, gross_ctc_pa=600000,
            basic_pm=25000, basic_pa=300000,
            net_salary_pm=45000, net_salary_pa=540000,
            ctc_pm=50000, ctc_pa=600000,
        )
        self.user = User.objects.create_superuser("salexp_admin", "a@test.com", "pass12345")
        self.client = Client()
        self.client.force_login(self.user, backend="django.contrib.auth.backends.ModelBackend")

    def test_export_url_is_reachable_and_returns_an_excel_file(self):
        resp = self.client.get(reverse("salary-master-export-excel"))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(
            resp["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def test_export_contains_the_salary_master_row_not_the_employee_list(self):
        resp = self.client.get(reverse("salary-master-export-excel"))
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        self.assertEqual(ws.title, "Salary Master")

        header = [c.value for c in ws[1]]
        self.assertIn("Gross CTC (PM)", header)
        self.assertIn("Basic (PM)", header)

        data_row = [c.value for c in ws[2]]
        row = dict(zip(header, data_row))
        self.assertEqual(row["Employee Code"], "SALEXP1")
        self.assertEqual(row["Gross CTC (PM)"], 50000)
        self.assertEqual(row["Basic (PM)"], 25000)

    def test_the_page_links_to_the_dedicated_export_endpoint_not_the_employee_export(self):
        resp = self.client.get(reverse("create-salary"))
        self.assertContains(resp, "downloadSalaryMaster")
        self.assertContains(resp, reverse("salary-master-export-excel"))


class SalaryIncrementExportTest(TestCase):
    def setUp(self):
        self.employee = make_employee("SALEXP2")
        self.active_inc = SalaryIncrement.objects.create(
            employee=self.employee,
            effective_date=date(2025, 1, 1),
            is_active=True,
            is_processed=True,
            change_set={
                "reason": "Annual raise",
                "flags": {"pf_deducted": True, "esic_applicable": False, "gratuity_applicable": True},
                "monthly": {"gross_ctc": 55000, "basic": 27500, "hra": 0, "net_salary": 50000, "ctc": 55000},
                "annual": {"gross_ctc": 660000, "basic": 330000, "hra": 0, "net_salary": 600000, "ctc": 660000},
            },
        )
        self.archived_inc = SalaryIncrement.objects.create(
            employee=self.employee,
            effective_date=date(2024, 1, 1),
            is_active=False,
            is_processed=False,
            change_set={"reason": "Old draft", "flags": {}, "monthly": {}, "annual": {}},
        )
        self.user = User.objects.create_superuser("salinc_admin", "b@test.com", "pass12345")
        self.client = Client()
        self.client.force_login(self.user, backend="django.contrib.auth.backends.ModelBackend")

    def test_export_url_is_reachable_and_returns_an_excel_file(self):
        resp = self.client.get(reverse("salary-increment-export-excel"))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(
            resp["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def test_default_export_only_includes_active_increments(self):
        resp = self.client.get(reverse("salary-increment-export-excel"))
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(rows), 1)
        header = [c.value for c in ws[1]]
        row = dict(zip(header, rows[0]))
        self.assertEqual(row["Reason"], "Annual raise")
        self.assertEqual(row["Status"], "Applied")
        self.assertEqual(row["Gross CTC (PM)"], 55000)

    def test_show_all_includes_the_archived_increment_too(self):
        resp = self.client.get(reverse("salary-increment-export-excel"), {"show": "all"})
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(rows), 2)

    def test_the_page_links_to_the_dedicated_export_endpoint_not_a_missing_function(self):
        resp = self.client.get(reverse("salary_increment"))
        self.assertContains(resp, "function exportIncrements")
        self.assertContains(resp, reverse("salary-increment-export-excel"))
